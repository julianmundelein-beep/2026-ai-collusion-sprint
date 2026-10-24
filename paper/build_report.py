#!/usr/bin/env python3
"""Build paper/report.docx on the official template from paper/report.md.

    python3 paper/build_report.py [--template "hackathon/official/Apart Research hackathon submission template.docx"]
                                  [--md paper/report.md] [--out paper/report.docx]

The template is copied and filled in place: title, author and affiliation in the title block, the abstract in the
abstract box, then every section from the markdown mirror with the template's own paragraph formats (Heading 2,
Heading 3, body text, numbered and bulleted lists). The italic guidance text and the yellow info box are removed.

Mirror syntax (paper/report.md):
  # Title · AUTHOR: name · AFFILIATION: text · ## Heading 2 · ### Heading 3 · paragraphs · "- " bullets ·
  "1. " numbered items · GFM tables (a line starting with "Table N." right before a table is its caption) ·
  ![caption](path) figures · ``` code blocks · **bold**, *italic*, `code` inline · lines starting with ">" are
  mirror notes and are not rendered · "[F13]" FACTS ids are stripped, any bracket with "F??" renders as "[F??]".
Needs python-docx (the paper venv or `pip install python-docx`).
"""
from __future__ import annotations

import argparse
import copy
import pathlib
import re
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Old Standard TT"
FACTS_ID = re.compile(r"\s*\[F\d{2}(?:\s*,\s*F\d{2})*\]")
FACTS_OPEN = re.compile(r"\[F\?\?[^\]]*\]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


# ---------------------------------------------------------------- markdown mirror

def parse_md(text: str) -> dict:
    """Return {'title', 'author', 'affiliation', 'abstract': [blocks], 'body': [blocks]}."""
    lines = text.splitlines()
    doc = {"title": "", "author": "", "affiliation": "", "abstract": [], "body": []}
    target = doc["body"]
    i = 0
    para: list[str] = []

    def flush():
        nonlocal para
        if para:
            target.append(("p", " ".join(s.strip() for s in para)))
            para = []

    while i < len(lines):
        ln = lines[i]
        s = ln.rstrip()
        if s.startswith("# ") and not doc["title"]:
            doc["title"] = s[2:].strip(); i += 1; continue
        if s.startswith("AUTHOR:"):
            doc["author"] = s.split(":", 1)[1].strip(); i += 1; continue
        if s.startswith("AFFILIATION:"):
            doc["affiliation"] = s.split(":", 1)[1].strip(); i += 1; continue
        if s.startswith(">"):
            flush(); i += 1; continue
        if not s.strip():
            flush(); i += 1; continue
        if s.startswith("## "):
            flush()
            head = s[3:].strip()
            if head.lower() == "abstract":
                target = doc["abstract"]
            else:
                if target is doc["abstract"]:
                    target = doc["body"]
                target.append(("h2", head))
            i += 1; continue
        if s.startswith("### "):
            flush(); target.append(("h3", s[4:].strip())); i += 1; continue
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            target.append(("code", code)); i += 1; continue
        m = re.match(r"!\[(.*)\]\((.*)\)$", s.strip())
        if m:
            flush(); target.append(("figure", m.group(2), m.group(1))); i += 1; continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            caption = None
            if target and target[-1][0] == "p" and re.match(r"Table \d+\.", target[-1][1]):
                caption = target.pop()[1]
            target.append(("table", rows, caption)); continue
        if re.match(r"- ", s):
            flush(); target.append(("bullet", s[2:].strip())); i += 1; continue
        if re.match(r"\d+\. ", s):
            flush(); target.append(("num", re.sub(r"^\d+\. ", "", s).strip())); i += 1; continue
        para.append(s); i += 1
    flush()
    return doc


def clean(text: str) -> str:
    text = FACTS_OPEN.sub("[F??]", text)
    text = FACTS_ID.sub("", text)
    return text


# ---------------------------------------------------------------- docx helpers

def run_props(bold=False, italic=False, mono=False, size=None, color=None):
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fam = "Courier New" if mono else FONT
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), fam)
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if color:
        c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    if size:
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag); e.set(qn("w:val"), str(int(size * 2))); rpr.append(e)
    return rpr


def make_run(text: str, **kw):
    r = OxmlElement("w:r")
    r.append(run_props(**kw))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_inline_runs(p, text: str, size=None, base_bold=False):
    """Append runs for **bold**, *italic*, `code` markup."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.append(make_run(part[2:-2], bold=True, size=size))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.append(make_run(part[1:-1], italic=True, size=size, bold=base_bold))
        elif part.startswith("`") and part.endswith("`"):
            p.append(make_run(part[1:-1], mono=True, size=(size or 11) - 1, bold=base_bold))
        else:
            p.append(make_run(part, size=size, bold=base_bold))


def strip_runs(p):
    for child in list(p):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            if child.tag == qn("w:r") and child.find(qn("w:footnoteReference")) is not None:
                continue  # keep the sprint footnote on the title
            p.remove(child)


def clone(proto):
    p = copy.deepcopy(proto)
    strip_runs(p)
    return p


def para_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


class Builder:
    def __init__(self, document, protos: dict):
        self.d = document
        self.body = document.element.body
        self.sect = self.body.find(qn("w:sectPr"))
        self.protos = protos

    def append(self, el):
        self.sect.addprevious(el)

    def paragraph(self, text, proto="normal", size=None, align=None, space_after=None, keep_next=False):
        p = clone(self.protos[proto])
        if proto == "normal" and size:
            pass
        add_inline_runs(p, clean(text), size=size)
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
        if align:
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc"); ppr.append(jc)
            jc.set(qn("w:val"), align)
        if space_after is not None:
            sp = ppr.find(qn("w:spacing"))
            if sp is None:
                sp = OxmlElement("w:spacing"); ppr.insert(0, sp)
            sp.set(qn("w:after"), str(space_after))
        if keep_next:
            kn = OxmlElement("w:keepNext"); ppr.append(kn)
        self.append(p)
        return p

    def code(self, lines):
        for ln in lines:
            p = clone(self.protos["normal"])
            ppr = p.find(qn("w:pPr"))
            sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), "0"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto")
            ppr.insert(0, sp)
            p.append(make_run(ln if ln else " ", mono=True, size=7.5))
            self.append(p)
        self.paragraph("", space_after=60)

    def figure(self, path: pathlib.Path, caption: str):
        if path.exists():
            self.d.add_picture(str(path), width=Inches(6.5))
            pic_p = self.body.findall(qn("w:p"))[-1]
            # add_picture inserts before sectPr already; centre it
            ppr = pic_p.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr"); pic_p.insert(0, ppr)
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc)
        else:
            self.paragraph(f"[Figure missing: {path}]", align="center")
        self.paragraph(caption, size=9.5, space_after=240)

    def table(self, rows: list[list[str]], caption: str | None):
        if caption:
            self.paragraph(caption, size=9.5, space_after=80)
        ncols = max(len(r) for r in rows)
        t = self.d.add_table(rows=len(rows), cols=ncols)
        tbl = t._tbl
        tblpr = tbl.find(qn("w:tblPr"))
        for child in list(tblpr):
            if child.tag in (qn("w:tblStyle"),):
                tblpr.remove(child)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "bottom", "insideH"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "808080")
            borders.append(b)
        for side in ("left", "right", "insideV"):
            b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), "nil"); borders.append(b)
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
        look = tblpr.find(qn("w:tblLook"))
        if look is not None:  # schema order: tblW, jc, tblInd, tblBorders, tblLayout, tblLook
            look.addprevious(borders); look.addprevious(layout)
        else:
            tblpr.append(borders); tblpr.append(layout)
        tw = tblpr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblpr.append(tw)
        tw.set(qn("w:w"), "9360"); tw.set(qn("w:type"), "dxa")
        # column widths by the longest cell text, bounded
        lens = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncols)]
        lens = [min(max(l, 4), 28) for l in lens]
        total = sum(lens)
        widths = [int(9360 * l / total) for l in lens]
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        size = 7
        for ri, row in enumerate(rows):
            tr = tbl.findall(qn("w:tr"))[ri]
            if ri == 0:
                trpr = tr.find(qn("w:trPr"))
                if trpr is None:
                    trpr = OxmlElement("w:trPr"); tr.insert(0, trpr)
                th = OxmlElement("w:tblHeader"); trpr.append(th)
            for ci in range(ncols):
                cell = t.cell(ri, ci)
                tc = cell._tc
                tcpr = tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW"); tcpr.insert(0, tcw)
                tcw.set(qn("w:w"), str(widths[ci])); tcw.set(qn("w:type"), "dxa")
                mar = OxmlElement("w:tcMar")
                for side, v in (("top", 15), ("bottom", 15), ("left", 50), ("right", 50)):
                    e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
                tcpr.append(mar)
                p = cell.paragraphs[0]._p
                ppr = p.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
                sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), "0"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto")
                ppr.append(sp)
                text = clean(row[ci]) if ci < len(row) else ""
                add_inline_runs(p, text, size=size, base_bold=(ri == 0))
        self.paragraph("", space_after=60)


def fill_title_block(document, title: str, author: str, affiliation: str):
    body = document.element.body
    for p in body.iter(qn("w:p")):
        if para_text(p) == "PROJECT TITLE":
            for r in p.findall(qn("w:r")):
                t = r.find(qn("w:t"))
                if t is not None and t.text == "PROJECT TITLE":
                    t.text = title
    author_ps = [p for p in body.iter(qn("w:p")) if para_text(p).startswith("Author name")]
    for k, p in enumerate(author_ps):
        for r in p.findall(qn("w:r")):
            p.remove(r)
        if k == 1:  # middle cell of the first row: a single author sits centred
            r = OxmlElement("w:r"); r.append(run_props(bold=True))
            t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = author
            br = OxmlElement("w:br"); br.set(qn("w:type"), "textWrapping")
            t2 = OxmlElement("w:t"); t2.set(qn("xml:space"), "preserve"); t2.text = affiliation
            r.append(t); r.append(br)
            r2 = OxmlElement("w:r"); r2.append(run_props()); r2.append(t2)
            p.append(r); p.append(r2)
    # drop the second author row (authors 4 to 6)
    for p in author_ps[3:]:
        tr = p.getparent().getparent()
        if tr.tag == qn("w:tr") and tr.getparent() is not None:
            tr.getparent().remove(tr)
            break


def fill_abstract(document, blocks):
    body = document.element.body
    for p in body.iter(qn("w:p")):
        if para_text(p).startswith("Summarize your project"):
            strip_runs(p)
            ppr = p.find(qn("w:pPr"))
            rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
            if rpr is not None:
                for u in rpr.findall(qn("w:u")):
                    rpr.remove(u)
            text = " ".join(b[1] for b in blocks if b[0] == "p")
            add_inline_runs(p, clean(text))
            return
    raise SystemExit("abstract box not found in the template")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    root = pathlib.Path(__file__).resolve().parent.parent
    ap.add_argument("--template", type=pathlib.Path,
                    default=root / "hackathon/official/Apart Research hackathon submission template.docx")
    ap.add_argument("--md", type=pathlib.Path, default=root / "paper/report.md")
    ap.add_argument("--out", type=pathlib.Path, default=root / "paper/report.docx")
    args = ap.parse_args(argv)

    md = parse_md(args.md.read_text(encoding="utf8"))
    shutil.copyfile(args.template, args.out)
    d = docx.Document(str(args.out))
    body = d.element.body
    top = [c for c in body if c.tag != qn("w:sectPr")]

    # prototypes from the guidance text, cloned before it is removed
    protos = {}
    for p in body.iter(qn("w:p")):
        t = para_text(p)
        if t == "1. Introduction":
            protos["h2"] = copy.deepcopy(p)
        elif t == "Limitations":
            protos["h3"] = copy.deepcopy(p)
        elif t.startswith("What problem are you addressing"):
            protos["normal"] = copy.deepcopy(p)
        elif t.startswith("[First contribution"):
            protos["num"] = copy.deepcopy(p)
        elif t == "[Reference 1]":
            protos["ref"] = copy.deepcopy(p)
        elif t.startswith("Code repository"):
            protos["bullet"] = copy.deepcopy(p)
    missing = {"h2", "h3", "normal", "num", "ref", "bullet"} - set(protos)
    if missing:
        raise SystemExit(f"template prototypes not found: {missing}")
    for key in ("h2", "h3"):  # headings: drop Google bookmarks
        for bm in protos[key].findall(qn("w:bookmarkStart")) + protos[key].findall(qn("w:bookmarkEnd")):
            protos[key].remove(bm)

    fill_title_block(d, clean(md["title"]), md["author"], md["affiliation"])
    fill_abstract(d, md["abstract"])

    # remove the info box and all guidance after the title block (index 1); keep the spacer paragraph at 2
    for el in top[3:]:
        body.remove(el)

    b = Builder(d, protos)
    ref_mode = False
    for blk in md["body"]:
        kind = blk[0]
        if kind == "h2":
            ref_mode = blk[1].strip().lower() == "references"
            b.paragraph(blk[1], proto="h2")
        elif kind == "h3":
            b.paragraph(blk[1], proto="h3")
        elif kind == "p":
            b.paragraph(blk[1])
        elif kind == "bullet":
            b.paragraph(blk[1], proto="bullet")
        elif kind == "num":
            b.paragraph(blk[1], proto="ref" if ref_mode else "num")
        elif kind == "code":
            b.code(blk[1])
        elif kind == "figure":
            b.figure((args.md.parent / blk[1]).resolve(), blk[2])
        elif kind == "table":
            b.table(blk[1], blk[2])
    d.save(str(args.out))
    words = sum(len(clean(x[1]).split()) for x in md["body"] if x[0] in ("p", "bullet", "num"))
    abstract_words = sum(len(clean(x[1]).split()) for x in md["abstract"] if x[0] == "p")
    print(f"wrote {args.out}: abstract {abstract_words} words, body text {words} words (tables, captions and code excluded)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
